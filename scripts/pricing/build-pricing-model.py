"""
Pricing and support model (draft for owner review; annual pricing approval is on hold).
Reworks Starter and Growth around a defined support allowance, paid onboarding and separately priced extra support,
tests viability at $40 per support hour, and sets out the pilot subsidy with a proposed spending cap.
All inputs are blue on the Assumptions and Pilot sheets. Usage: python3 build-pricing-model.py out.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = 'Arial'
BLUE = Font(name=F, color='0000FF')
BLACK = Font(name=F)
GREEN = Font(name=F, color='008000')
BOLD = Font(name=F, bold=True)
NOTE = Font(name=F, size=10, color='475569')
HEAD = Font(name=F, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='0E7490')
KEY_FILL = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='CBD5E1')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0.00;($#,##0.00);-'
USD0 = '$#,##0;($#,##0);-'
PCT = '0.0%;(0.0%);-'
NUM = '#,##0.00;(#,##0.00);-'
PLANS = ['Starter', 'Growth', 'Business', 'Enterprise']
COL = {p: chr(ord('B') + i) for i, p in enumerate(PLANS)}


def head(ws, row, labels, height=None):
    for j, h in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill, c.border = HEAD, HEAD_FILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if height:
        ws.row_dimensions[row].height = height


def inputs_sheet(wb):
    a = wb.create_sheet('Assumptions')
    a.column_dimensions['A'].width = 52
    for c in 'BCDE':
        a.column_dimensions[c].width = 13
    a.column_dimensions['F'].width = 78
    a.cell(row=1, column=1, value='Assumptions — change only the blue numbers').font = Font(name=F, bold=True, size=13)
    head(a, 2, ['Per plan', *PLANS, 'Note / source'])
    plan_rows = [
        ('Monthly price ($)', [29, 79, 199, 499], USD0, 'BRD v1.1 proposed prices (catalogue version 1)', False),
        ('Seat limit', [3, 10, 25, 50], '0', 'BRD v1.1 plan limits', False),
        ('Share of seats in use', [0.7, 0.7, 0.7, 0.7], PCT, 'Assumption', False),
        ('Included automation runs / month', [1000, 5000, 25000, 50000], '#,##0', 'BRD v1.1; Enterprise is custom — 50,000 assumed', False),
        ('Support allowance included (hours / month)', [0.25, 0.75, 3, 6], NUM,
         'Proposed: Starter 15 min (email), Growth 45 min (email and chat); Business and Enterprise unchanged', True),
        ('Support customers expect (hours / month)', [0.5, 1.5, 3, 6], NUM, 'Same demand as the previous model, so the comparison is fair', True),
        ('Share of time above the allowance bought as extra support', [0.4, 0.4, 0.4, 0.4], PCT,
         'Assumption: the rest is handled by the help centre or not requested', False),
        ('Onboarding fee, one-off ($)', [99, 199, 399, 799], USD0,
         'Proposed: Starter optional assisted set-up; Growth and above include paid onboarding', True),
        ('Share of new companies paying for onboarding', [0.3, 1.0, 1.0, 1.0], PCT, 'Starter is optional; Growth and above required', False),
        ('Onboarding time when paid (hours)', [1.5, 3, 6, 12], NUM, 'Assumption: guided set-up call, templates, team and roles', False),
        ('Onboarding time otherwise (hours)', [0.5, 0.5, 0, 0], NUM, 'Assumption: self-service with one short check-in', False),
        ('Companies at growth scale', [40, 35, 20, 5], '0', 'Assumption: plan mix of 100 companies', False),
    ]
    R = {}
    for i, (label, vals, fmt, note, key) in enumerate(plan_rows, start=3):
        a.cell(row=i, column=1, value=label).font = BLACK
        for j, v in enumerate(vals, start=2):
            c = a.cell(row=i, column=j, value=v)
            c.font, c.number_format, c.border = BLUE, fmt, BORDER
            if key:
                c.fill = KEY_FILL
        a.cell(row=i, column=6, value=note).font = NOTE
        R[label] = i
    g0 = 3 + len(plan_rows) + 1
    head(a, g0, ['Shared inputs', 'Value', '', '', '', 'Note / source'])
    shared = [
        ('Hosting at growth scale ($ / month)', 1554, USD0, 'Provisional AWS Frankfurt estimate excl. email (docs/readiness/hosting-comparison.md)', False),
        ('Share of hosting split evenly per company', 0.3, PCT, 'Assumption: the rest is split by seats in use', False),
        ('Emails per seat in use / month', 20, '0', 'Assumption', False),
        ('Email cost per 1,000 ($)', 0.10, USD, 'Amazon SES Frankfurt list price', False),
        ('File storage per seat in use (GB)', 1, NUM, 'Assumption', False),
        ('Storage cost per GB-month incl. backup copy ($)', 0.05, USD, 'Assumption', False),
        ('Compute cost per 1,000 automation runs ($)', 0.05, USD, 'Assumption', False),
        ('Share of included runs actually used', 0.5, PCT, 'Assumption', False),
        ('Support cost per hour, fully loaded ($)', 40, USD0, 'Owner test value, 20 Oct 2026', True),
        ('Extra support price per hour ($)', 60, USD0, 'Proposed: sold in 30-minute steps, separate from the plan price', True),
        ('Average customer lifetime (months)', 24, '0', 'Assumption: spreads one-off onboarding over the months it serves', False),
        ('Card processing — percent of charge', 0.03, PCT, 'Typical; varies by processor and country', False),
        ('Card processing — fixed fee per charge ($)', 0.30, USD, 'Typical', False),
        ('Annual price multiple (× monthly)', 10, '0', 'Draft, unpublished; approval on hold', False),
    ]
    G = {}
    for k, (label, v, fmt, note, key) in enumerate(shared, start=g0 + 1):
        a.cell(row=k, column=1, value=label).font = BLACK
        c = a.cell(row=k, column=2, value=v)
        c.font, c.number_format, c.border = BLUE, fmt, BORDER
        if key:
            c.fill = KEY_FILL
        a.cell(row=k, column=6, value=note).font = NOTE
        G[label] = f'Assumptions!$B${k}'
    leg = g0 + len(shared) + 2
    a.cell(row=leg, column=1, value='Blue = input you can change · yellow = biggest effect on the answer · black = calculated · green = taken from another sheet.').font = Font(name=F, italic=True, color='475569')
    A = lambda label, p: f'Assumptions!${COL[p]}${R[label]}'
    return A, G


def margin_sheet(wb, A, G):
    m = wb.create_sheet('Margin')
    m.cell(row=1, column=1, value='Growth scale — 100 companies, per company per month (support costed at the hourly rate on Assumptions)').font = Font(name=F, bold=True, size=12)
    heads = ['Plan', 'Monthly price', 'Annual, per month', 'Hosting', 'Included usage', 'Included support (cost)',
             'Extra support (revenue)', 'Extra support (cost)', 'Onboarding fee (per month)', 'Onboarding time (cost per month)',
             'Card fees, monthly billing', 'Card fees, annual billing', 'Margin, monthly billing', 'Margin %, monthly',
             'Margin, annual billing', 'Margin %, annual', 'Support hours carried per month']
    head(m, 2, heads, 58)
    for j, w in enumerate([14] + [12] * 16, start=1):
        m.column_dimensions[chr(64 + j)].width = w
    seats_used = lambda p: f'{A("Seat limit", p)}*{A("Share of seats in use", p)}'
    cnt = lambda p: A('Companies at growth scale', p)
    total_c = '(' + '+'.join(cnt(p) for p in PLANS) + ')'
    total_s = '(' + '+'.join(f'{cnt(p)}*{seats_used(p)}' for p in PLANS) + ')'
    cost, price, life = G['Support cost per hour, fully loaded ($)'], G['Extra support price per hour ($)'], G['Average customer lifetime (months)']
    pct, fix = G['Card processing — percent of charge'], G['Card processing — fixed fee per charge ($)']
    first = 3
    for i, p in enumerate(PLANS):
        r = first + i
        allow, demand, uptake = A('Support allowance included (hours / month)', p), A('Support customers expect (hours / month)', p), A('Share of time above the allowance bought as extra support', p)
        excess = f'MAX(0,{demand}-{allow})*{uptake}'
        share_paid = A('Share of new companies paying for onboarding', p)
        onb_hours = f'({share_paid}*{A("Onboarding time when paid (hours)", p)}+(1-{share_paid})*{A("Onboarding time otherwise (hours)", p)})'
        share = G['Share of hosting split evenly per company']
        f = {
            1: p,
            2: f'={A("Monthly price ($)", p)}',
            3: f'=B{r}*{G["Annual price multiple (× monthly)"]}/12',
            4: f'=IF({total_c}=0,0,{G["Hosting at growth scale ($ / month)"]}*({share}/{total_c}+(1-{share})*{seats_used(p)}/{total_s}))',
            5: (f'={seats_used(p)}*{G["Emails per seat in use / month"]}/1000*{G["Email cost per 1,000 ($)"]}'
                f'+{seats_used(p)}*{G["File storage per seat in use (GB)"]}*{G["Storage cost per GB-month incl. backup copy ($)"]}'
                f'+{A("Included automation runs / month", p)}*{G["Share of included runs actually used"]}/1000*{G["Compute cost per 1,000 automation runs ($)"]}'),
            6: f'=MIN({allow},{demand})*{cost}',
            7: f'={excess}*{price}',
            8: f'={excess}*{cost}',
            9: f'={A("Onboarding fee, one-off ($)", p)}*{share_paid}/{life}',
            10: f'={onb_hours}*{cost}/{life}',
            11: f'=(B{r}+G{r}+I{r})*{pct}+{fix}',
            12: f'=(C{r}*12*{pct}+{fix})/12+(G{r}+I{r})*{pct}',
            13: f'=B{r}+G{r}+I{r}-D{r}-E{r}-F{r}-H{r}-J{r}-K{r}',
            14: f'=IF(B{r}+G{r}+I{r}=0,0,M{r}/(B{r}+G{r}+I{r}))',
            15: f'=C{r}+G{r}+I{r}-D{r}-E{r}-F{r}-H{r}-J{r}-L{r}',
            16: f'=IF(C{r}+G{r}+I{r}=0,0,O{r}/(C{r}+G{r}+I{r}))',
            17: f'=MIN({allow},{demand})+{excess}+{onb_hours}/{life}',
        }
        for j, v in f.items():
            c = m.cell(row=r, column=j, value=v)
            c.border = BORDER
            c.font = BOLD if j == 1 else (GREEN if j == 2 else BLACK)
            if j > 1:
                c.number_format = PCT if j in (14, 16) else (NUM if j == 17 else USD)
    tr = first + len(PLANS)
    mix = lambda L: '+'.join(f'{cnt(p)}*{L}{first + i}' for i, p in enumerate(PLANS))
    m.cell(row=tr, column=1, value='Whole mix').font = BOLD
    for j, L in [(2, 'B'), (3, 'C'), (7, 'G'), (9, 'I'), (13, 'M'), (15, 'O')]:
        m.cell(row=tr, column=j, value=f'={mix(L)}')
    m.cell(row=tr, column=14, value=f'=IF(B{tr}+G{tr}+I{tr}=0,0,M{tr}/(B{tr}+G{tr}+I{tr}))')
    m.cell(row=tr, column=16, value=f'=IF(C{tr}+G{tr}+I{tr}=0,0,O{tr}/(C{tr}+G{tr}+I{tr}))')
    for j in range(1, 18):
        c = m.cell(row=tr, column=j)
        c.border, c.font = BORDER, BOLD
        if j > 1:
            c.number_format = PCT if j in (14, 16) else USD0
    m.cell(row=tr + 1, column=1, value='Whole mix = monthly totals across all 100 companies. Revenue includes extra support and onboarding fees spread over the customer lifetime.').font = Font(name=F, size=9, italic=True, color='475569')

    # Before and after, at the same support cost
    b0 = tr + 3
    m.cell(row=b0, column=1, value='Before and after the rework — annual billing, at the support cost on Assumptions').font = Font(name=F, bold=True, size=12)
    head(m, b0 + 1, ['Plan', 'Before: all expected support free', 'Before: margin %', 'After: margin per month', 'After: margin %', 'Change per month'], 44)
    for i, p in enumerate(PLANS):
        r, src = b0 + 2 + i, first + i
        before = f'C{src}-D{src}-E{src}-{A("Support customers expect (hours / month)", p)}*{cost}-(C{src}*12*{pct}+{fix})/12'
        vals = {1: p, 2: f'={before}', 3: f'=IF(C{src}=0,0,B{r}/C{src})', 4: f'=O{src}', 5: f'=P{src}', 6: f'=D{r}-B{r}'}
        for j, v in vals.items():
            c = m.cell(row=r, column=j, value=v)
            c.border = BORDER
            c.font = BOLD if j == 1 else BLACK
            if j > 1:
                c.number_format = PCT if j in (3, 5) else USD
    # Sensitivity: margin is linear in the support cost per hour
    s0 = b0 + 2 + len(PLANS) + 2
    m.cell(row=s0, column=1, value='Annual-billing margin per company per month at different support costs per hour (after the rework)').font = Font(name=F, bold=True, size=12)
    rates = [20, 30, 40, 50, 60]
    head(m, s0 + 1, ['Plan', *[f'${x} / hour' for x in rates]], 30)
    for i, p in enumerate(PLANS):
        r, src = s0 + 2 + i, first + i
        m.cell(row=r, column=1, value=p).font = BOLD
        m.cell(row=r, column=1).border = BORDER
        for j, x in enumerate(rates, start=2):
            c = m.cell(row=r, column=j, value=f'=O{src}+({cost}-{x})*Q{src}')
            c.number_format, c.border, c.font = USD, BORDER, BLACK
    m.cell(row=s0 + 2 + len(PLANS), column=1, value='Extra support is sold at the price on Assumptions ($60 an hour). If support costs rise above that price, extra support no longer covers its own time.').font = Font(name=F, size=9, italic=True, color='475569')
    m.freeze_panes = 'B3'
    return m


def pilot_sheet(wb, G):
    p = wb.create_sheet('Pilot subsidy')
    p.column_dimensions['A'].width = 58
    p.column_dimensions['B'].width = 16
    p.column_dimensions['C'].width = 16
    p.column_dimensions['D'].width = 70
    p.cell(row=1, column=1, value='Pilot subsidy — what the pilot costs Meronatic, and a proposed spending cap').font = Font(name=F, bold=True, size=13)
    p.cell(row=2, column=1, value='Proposed pilot terms: no ConversaFlow fee; Meta message fees paid by each pilot company directly.').font = NOTE
    head(p, 4, ['Input', '3 companies', '5 companies', 'Note'])
    inputs = [
        ('Pilot companies', [3, 5], '0', 'One per sector; up to five if you choose more'),
        ('Pilot length (days)', [90, 90], '0', 'Same as the system default for a pilot enrolment'),
        ('Onboarding time per company (hours)', [8, 8], NUM, 'Kick-off, workspace set-up, templates, training, closing summary'),
        ('Weeks of close support', [2, 2], '0', 'From the onboarding plan in the preparation pack'),
        ('Support hours per week, close support', [4, 4], NUM, 'Assumption: daily check-ins plus the weekly review'),
        ('Support hours per week, afterwards', [1.5, 1.5], NUM, 'Assumption: weekly 30-minute review plus questions'),
        ('Hosting during the pilot ($ / month)', [441, 441], USD0, 'Provisional AWS Frankfurt pilot estimate excl. email (production and staging)'),
        ('Other running costs per company ($ / month)', [5, 5], USD0, 'Email, storage, logs: assumption, small'),
        ('ConversaFlow fee per company during the pilot ($ / month)', [0, 0], USD0, 'Proposed pilot terms: no fee'),
        ('Contingency on top of the estimate', [0.15, 0.15], PCT, 'For extra support time or a longer pilot'),
    ]
    R = {}
    for i, (label, vals, fmt, note) in enumerate(inputs, start=5):
        p.cell(row=i, column=1, value=label).font = BLACK
        for j, v in enumerate(vals, start=2):
            c = p.cell(row=i, column=j, value=v)
            c.font, c.number_format, c.border = BLUE, fmt, BORDER
        p.cell(row=i, column=4, value=note).font = NOTE
        R[label] = i
    r0 = 5 + len(inputs)
    p.cell(row=r0, column=1, value='Support cost per hour ($)').font = BLACK
    for col in 'BC':
        c = p[f'{col}{r0}']
        c.value, c.font, c.number_format, c.border = f'={G["Support cost per hour, fully loaded ($)"]}', GREEN, USD0, BORDER
    p.cell(row=r0, column=4, value='From Assumptions').font = NOTE
    R['cost'] = r0

    c0 = r0 + 2
    head(p, c0, ['Calculation', '3 companies', '5 companies', 'How it is worked out'])
    I = lambda label, col: f'{col}{R[label]}'
    calc = [
        ('Pilot length (months)', lambda c: f'={I("Pilot length (days)", c)}/30.4', NUM, 'Days ÷ 30.4'),
        ('Support hours per company', lambda c: (
            f'={I("Onboarding time per company (hours)", c)}+{I("Weeks of close support", c)}*{I("Support hours per week, close support", c)}'
            f'+MAX(0,{I("Pilot length (days)", c)}/7-{I("Weeks of close support", c)})*{I("Support hours per week, afterwards", c)}'), NUM,
         'Onboarding + close-support weeks + the remaining weeks'),
        ('Support time ($)', lambda c: f'={I("Pilot companies", c)}*{{h}}*{c}{R["cost"]}', USD0, 'Companies × hours × hourly cost'),
        ('Hosting ($)', lambda c: f'={{months}}*{I("Hosting during the pilot ($ / month)", c)}', USD0, 'Months × monthly hosting'),
        ('Other running costs ($)', lambda c: f'={I("Pilot companies", c)}*{{months}}*{I("Other running costs per company ($ / month)", c)}', USD0, 'Companies × months × cost'),
        ('Total cost ($)', lambda c: '={sum}', USD0, 'Support + hosting + other'),
        ('Less pilot fees received ($)', lambda c: f'={I("Pilot companies", c)}*{{months}}*{I("ConversaFlow fee per company during the pilot ($ / month)", c)}', USD0, 'Zero under the proposed terms'),
        ('Pilot subsidy ($)', lambda c: '={total}-{fees}', USD0, 'What Meronatic pays for the pilot'),
        ('Proposed spending cap ($)', lambda c: f'=CEILING({{subsidy}}*(1+{I("Contingency on top of the estimate", c)}),500)', USD0,
         'Subsidy plus contingency, rounded up to the next $500'),
    ]
    rows = {}
    for k, (label, fn, fmt, how) in enumerate(calc, start=c0 + 1):
        rows[label] = k
    for k, (label, fn, fmt, how) in enumerate(calc, start=c0 + 1):
        p.cell(row=k, column=1, value=label).font = BOLD if label in ('Pilot subsidy ($)', 'Proposed spending cap ($)') else BLACK
        for col in 'BC':
            v = fn(col).format(
                h=f'{col}{rows["Support hours per company"]}', months=f'{col}{rows["Pilot length (months)"]}',
                sum=f'{col}{rows["Support time ($)"]}+{col}{rows["Hosting ($)"]}+{col}{rows["Other running costs ($)"]}',
                total=f'{col}{rows["Total cost ($)"]}', fees=f'{col}{rows["Less pilot fees received ($)"]}',
                subsidy=f'{col}{rows["Pilot subsidy ($)"]}')
            c = p[f'{col}{k}']
            c.value, c.number_format, c.border = v, fmt, BORDER
            c.font = BOLD if label in ('Pilot subsidy ($)', 'Proposed spending cap ($)') else BLACK
            if label == 'Proposed spending cap ($)':
                c.fill = KEY_FILL
        p.cell(row=k, column=4, value=how).font = NOTE
    n = c0 + len(calc) + 2
    for t in [
        'Not included: the independent security review (quoted separately), legal drafting, Meta fees (paid by the pilot companies) and development time.',
        'Hosting is shared: it is the same whether the pilot has three or five companies, so more companies lower the subsidy per company.',
        'How the cap would work: spending is tracked monthly against the cap; reaching 80% triggers a review with you before more is committed.',
    ]:
        p.cell(row=n, column=1, value=t).font = NOTE
        n += 1


def offer_sheet(wb):
    o = wb.create_sheet('Support offer')
    o.column_dimensions['A'].width = 30
    for c in 'BCDE':
        o.column_dimensions[c].width = 34
    o.cell(row=1, column=1, value='Proposed support offer per plan (draft for your decision)').font = Font(name=F, bold=True, size=13)
    head(o, 3, ['', *PLANS], 22)
    rows = [
        ('Included support each month', '15 minutes by email, next business day', '45 minutes by email and chat, same business day', '3 hours, priority, named contact', '6 hours, priority, named contact'),
        ('Onboarding', 'Self-service guide; assisted set-up optional at $99', 'Paid onboarding $199, one-off (about 3 hours)', 'Paid onboarding $399 (about 6 hours)', 'Paid onboarding from $799 (about 12 hours)'),
        ('Extra support', '$60 per hour, in 30-minute steps', '$60 per hour, in 30-minute steps', '$60 per hour, in 30-minute steps', 'Included in the contract or $60 per hour'),
        ('Help centre and guides', 'Yes', 'Yes', 'Yes', 'Yes'),
        ('Annual billing (10 × monthly)', 'On hold until you approve', 'On hold until you approve', 'On hold until you approve', 'On hold until you approve'),
    ]
    for i, row in enumerate(rows, start=4):
        for j, v in enumerate(row, start=1):
            c = o.cell(row=i, column=j, value=v)
            c.border = BORDER
            c.font = BOLD if j == 1 else BLACK
            c.alignment = Alignment(wrap_text=True, vertical='top')
        o.row_dimensions[i].height = 34
    o.cell(row=10, column=1, value='All figures are proposals for your decision. Nothing is published or offered to customers.').font = NOTE


def readme(wb):
    rm = wb.active
    rm.title = 'Read me'
    rm.column_dimensions['A'].width = 120
    lines = [
        ('ConversaFlow — pricing and support model (draft)', Font(name=F, bold=True, size=16)),
        ('Status: annual pricing approval is on hold (owner, 20 Oct 2026). Nothing here is published or offered.', None),
        ('', None),
        ('What changed from the earlier model', Font(name=F, bold=True, size=12)),
        ('• Starter and Growth now include a defined support allowance (15 and 45 minutes a month) instead of absorbing all the support customers ask for.', None),
        ('• Onboarding is paid where it is needed (optional for Starter, included as a fee for Growth and above), and extra support is priced separately per hour.', None),
        ('• Support is costed at $40 an hour (your test value) and a sensitivity table shows $20 to $60.', None),
        ('• The Pilot subsidy sheet shows the total cost of a free pilot for three or five companies and a proposed spending cap.', None),
        ('', None),
        ('How to use', Font(name=F, bold=True, size=12)),
        ('• Change only blue numbers (Assumptions and Pilot subsidy). Yellow marks the inputs with the biggest effect.', None),
        ('• Margin compares annual and monthly billing per plan, before and after the rework, and across support costs.', None),
        ('', None),
        ('Excluded', Font(name=F, bold=True, size=12)),
        ('• Meta message fees (paid by customers directly), usage overages, AI credits (later phase), sales commission, taxes and development cost.', None),
        ('• Hosting figures are provisional estimates (docs/readiness/hosting-comparison.md); no hosting is selected or bought.', None),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = rm.cell(row=i, column=1, value=t)
        c.font = f or Font(name=F, size=11)
        c.alignment = Alignment(wrap_text=True, vertical='top')


def main(out):
    wb = Workbook()
    readme(wb)
    A, G = inputs_sheet(wb)
    margin_sheet(wb, A, G)
    pilot_sheet(wb, G)
    offer_sheet(wb)
    wb.move_sheet('Margin', offset=-2)
    wb.save(out)
    print(out)


if __name__ == '__main__':
    main(sys.argv[1])
