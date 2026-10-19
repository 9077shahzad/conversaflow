"""
Annual pricing margin model (draft for owner review). All inputs are on the Assumptions sheet (blue = editable).
Usage: python3 build-margin-model.py out.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = 'Arial'
BLUE = Font(name=F, color='0000FF')
BLACK = Font(name=F)
GREEN = Font(name=F, color='008000')
BOLD = Font(name=F, bold=True)
HEAD = Font(name=F, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='0E7490')
KEY_FILL = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='CBD5E1')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0.00;($#,##0.00);-'
USD0 = '$#,##0;($#,##0);-'
PCT = '0.0%;(0.0%);-'
NUM = '#,##0.0;(#,##0.0);-'

PLANS = ['Starter', 'Growth', 'Business', 'Enterprise']


def main(out):
    wb = Workbook()

    # ---------------- Read me ----------------
    rm = wb.active
    rm.title = 'Read me'
    rm.sheet_view.showGridLines = False
    text = [
        ('ConversaFlow — annual pricing margin model (draft)', Font(name=F, bold=True, size=16)),
        ('Purpose: show the expected gross margin per plan, on monthly and on annual billing (annual = 10 × monthly), after infrastructure, customer support and the usage included in each plan — before final pricing approval.', None),
        ('Status: draft for margin review. The annual prices are stored unpublished in the system; nothing is offered to customers.', None),
        ('', None),
        ('How to use', Font(name=F, bold=True, size=12)),
        ('• Change only the blue numbers on the Assumptions sheet. Yellow marks the assumptions with the biggest effect on margin.', None),
        ('• The Margin sheet recalculates. It shows two scales: a pilot (5 companies) and growth (100 companies).', None),
        ('', None),
        ('What is excluded (by owner decision, 19 Oct 2026)', Font(name=F, bold=True, size=12)),
        ('• Meta (WhatsApp) message fees — each customer pays Meta directly.', None),
        ('• Usage overages — billed separately when they occur.', None),
        ('• Not yet modelled: AI costs (Phase 4, to be sold as credits), sales commission, taxes, development cost.', None),
        ('', None),
        ('Where the cost figures come from', Font(name=F, bold=True, size=12)),
        ('• Hosting: EU hosting comparison of 19 Oct 2026 (AWS Frankfurt, list prices, estimates; see the Pre-Pilot Preparation Pack). Some AWS Frankfurt prices came from third-party mirrors and should be checked in the AWS calculator.', None),
        ('• Email: Amazon SES Frankfurt list price, $0.10 per 1,000 emails.', None),
        ('• Support time, support cost and included-usage costs: working assumptions, not sourced — replace with your own figures.', None),
    ]
    for i, (t, f) in enumerate(text, start=1):
        c = rm.cell(row=i, column=1, value=t)
        c.font = f or Font(name=F, size=11)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    rm.column_dimensions['A'].width = 120

    # ---------------- Assumptions ----------------
    a = wb.create_sheet('Assumptions')
    a.column_dimensions['A'].width = 46
    for col in 'BCDE':
        a.column_dimensions[col].width = 15
    a.column_dimensions['F'].width = 70

    def header(row, labels):
        for j, h in enumerate(labels, start=1):
            c = a.cell(row=row, column=j, value=h)
            c.font, c.fill, c.border = HEAD, HEAD_FILL, BORDER

    # Plan table (rows 3-?)
    header(2, ['Per plan', *PLANS, 'Note / source'])
    plan_rows = [
        ('Monthly price ($)', [29, 79, 199, 499], USD0, 'BRD v1.1 proposed prices (owner-approved catalogue, version 1)'),
        ('Seat limit', [3, 10, 25, 50], '0', 'BRD v1.1 plan limits'),
        ('Share of seats in use', [0.7, 0.7, 0.7, 0.7], PCT, 'Assumption: companies use 70% of their seat limit'),
        ('Included automation runs / month', [1000, 5000, 25000, 50000], '#,##0', 'BRD v1.1; Enterprise is custom — 50,000 assumed for costing'),
        ('Support hours per company / month', [0.5, 1.5, 3, 6], NUM, 'Assumption — replace with your support plan'),
        ('Companies at growth scale', [40, 35, 20, 5], '0', 'Assumption: plan mix of 100 companies'),
        ('Companies at pilot scale', [1, 2, 2, 0], '0', 'Assumption: 5 pilot companies'),
    ]
    rows = {}
    for i, (label, vals, fmt, note) in enumerate(plan_rows, start=3):
        a.cell(row=i, column=1, value=label).font = BLACK
        for j, v in enumerate(vals, start=2):
            c = a.cell(row=i, column=j, value=v)
            c.font, c.number_format, c.border = BLUE, fmt, BORDER
            if label.startswith('Support hours'):
                c.fill = KEY_FILL
        a.cell(row=i, column=6, value=note).font = Font(name=F, size=10, color='475569')
        rows[label] = i

    # Global inputs
    g0 = 3 + len(plan_rows) + 1
    header(g0, ['Shared inputs', 'Value', '', '', '', 'Note / source'])
    globals_ = [
        ('Hosting at growth scale ($ / month)', 1554, USD0, 'AWS Frankfurt growth estimate excl. email (docs/readiness/hosting-comparison.md): HA database, app and ClamAV containers, 1 TB files, WAF, backups, logs, support plan', True),
        ('Hosting at pilot scale ($ / month)', 441, USD0, 'AWS Frankfurt pilot estimate excl. email (docs/readiness/hosting-comparison.md): production + smaller staging', True),
        ('Emails per seat in use / month', 20, '0', 'Assumption: invitations, resets, notifications', False),
        ('Email cost per 1,000 ($)', 0.10, USD, 'Amazon SES Frankfurt list price', False),
        ('Share of hosting split evenly per company', 0.3, PCT, 'Assumption: rest is split by seats in use', False),
        ('File storage per seat in use (GB)', 1, NUM, 'Assumption', False),
        ('Storage cost per GB-month incl. backup copy ($)', 0.05, USD, 'Assumption: S3 Frankfurt about $0.0245/GB plus a second-region copy', False),
        ('Compute cost per 1,000 automation runs ($)', 0.05, USD, 'Assumption: worker and database time', False),
        ('Share of included runs actually used', 0.5, PCT, 'Assumption', False),
        ('Support cost per hour, fully loaded ($)', 20, USD, 'Assumption — e.g. a GCC-based support agent incl. overheads', True),
        ('Card processing — percent of charge', 0.03, PCT, 'Typical card fee; varies by processor and country', False),
        ('Card processing — fixed fee per charge ($)', 0.30, USD, 'Typical card fee', False),
        ('Annual price multiple (× monthly)', 10, '0', 'Owner decision 19 Oct 2026 (draft, unpublished)', False),
    ]
    G = {}
    for k, (label, v, fmt, note, key) in enumerate(globals_, start=g0 + 1):
        a.cell(row=k, column=1, value=label).font = BLACK
        c = a.cell(row=k, column=2, value=v)
        c.font, c.number_format, c.border = BLUE, fmt, BORDER
        if key:
            c.fill = KEY_FILL
        a.cell(row=k, column=6, value=note).font = Font(name=F, size=10, color='475569')
        G[label] = f'Assumptions!$B${k}'
    leg = g0 + len(globals_) + 2
    a.cell(row=leg, column=1, value='Legend: blue = input you can change · yellow = biggest effect on margin · black = calculated.').font = Font(name=F, italic=True, color='475569')

    col = {p: chr(ord('B') + i) for i, p in enumerate(PLANS)}
    A = lambda label, p: f'Assumptions!${col[p]}${rows[label]}'

    # ---------------- Margin ----------------
    m = wb.create_sheet('Margin')
    widths = [16, 12, 13, 13, 14, 12, 13, 13, 13, 13, 12, 13, 12]
    heads = ['Plan', 'Monthly price', 'Annual price (draft)', 'Annual, per month',
             'Hosting', 'Support', 'Included usage', 'Card fees (monthly billing)', 'Card fees (annual, per month)',
             'Margin, monthly billing', 'Margin %, monthly', 'Margin, annual billing (per month)', 'Margin %, annual']
    for j, w in enumerate(widths, start=1):
        m.column_dimensions[chr(64 + j)].width = w

    def block(start, title, scale_label, hosting_ref):
        m.cell(row=start, column=1, value=title).font = Font(name=F, bold=True, size=12)
        for j, h in enumerate(heads, start=1):
            c = m.cell(row=start + 1, column=j, value=h)
            c.font, c.fill, c.border = HEAD, HEAD_FILL, BORDER
            c.alignment = Alignment(wrap_text=True, vertical='center')
        m.row_dimensions[start + 1].height = 44
        cnt = lambda p: A(scale_label, p)
        seats_used = lambda p: f'{A("Seat limit", p)}*{A("Share of seats in use", p)}'
        total_companies = '(' + '+'.join(cnt(p) for p in PLANS) + ')'
        total_seats = '(' + '+'.join(f'{cnt(p)}*{seats_used(p)}' for p in PLANS) + ')'
        first = start + 2
        for i, p in enumerate(PLANS):
            r = first + i
            m.cell(row=r, column=1, value=p).font = BOLD
            m.cell(row=r, column=2, value=f'={A("Monthly price ($)", p)}').font = GREEN
            m.cell(row=r, column=3, value=f'=B{r}*{G["Annual price multiple (× monthly)"]}')
            m.cell(row=r, column=4, value=f'=C{r}/12')
            share = G['Share of hosting split evenly per company']
            m.cell(row=r, column=5, value=(
                f'=IF({total_companies}=0,0,{hosting_ref}*({share}/{total_companies}'
                f'+(1-{share})*{seats_used(p)}/{total_seats}))'))
            m.cell(row=r, column=6, value=f'={A("Support hours per company / month", p)}*{G["Support cost per hour, fully loaded ($)"]}')
            m.cell(row=r, column=7, value=(
                f'={seats_used(p)}*{G["Emails per seat in use / month"]}/1000*{G["Email cost per 1,000 ($)"]}'
                f'+{seats_used(p)}*{G["File storage per seat in use (GB)"]}*{G["Storage cost per GB-month incl. backup copy ($)"]}'
                f'+{A("Included automation runs / month", p)}*{G["Share of included runs actually used"]}/1000*{G["Compute cost per 1,000 automation runs ($)"]}'))
            pct, fix = G['Card processing — percent of charge'], G['Card processing — fixed fee per charge ($)']
            m.cell(row=r, column=8, value=f'=B{r}*{pct}+{fix}')
            m.cell(row=r, column=9, value=f'=(C{r}*{pct}+{fix})/12')
            m.cell(row=r, column=10, value=f'=B{r}-E{r}-F{r}-G{r}-H{r}')
            m.cell(row=r, column=11, value=f'=IF(B{r}=0,0,J{r}/B{r})')
            m.cell(row=r, column=12, value=f'=D{r}-E{r}-F{r}-G{r}-I{r}')
            m.cell(row=r, column=13, value=f'=IF(D{r}=0,0,L{r}/D{r})')
            for j in range(1, 14):
                c = m.cell(row=r, column=j)
                c.border = BORDER
                if j > 1:
                    c.number_format = PCT if j in (11, 13) else (USD0 if j in (2, 3) else USD)
                if j > 2 and c.font != GREEN:
                    c.font = BLACK
        # Weighted total (monthly gross profit across the mix)
        tr = first + len(PLANS)
        m.cell(row=tr, column=1, value='Whole mix').font = BOLD
        mix = lambda colL: '+'.join(f'{cnt(p)}*{colL}{first + i}' for i, p in enumerate(PLANS))
        m.cell(row=tr, column=2, value=f'={mix("B")}')
        m.cell(row=tr, column=4, value=f'={mix("D")}')
        m.cell(row=tr, column=10, value=f'={mix("J")}')
        m.cell(row=tr, column=11, value=f'=IF(B{tr}=0,0,J{tr}/B{tr})')
        m.cell(row=tr, column=12, value=f'={mix("L")}')
        m.cell(row=tr, column=13, value=f'=IF(D{tr}=0,0,L{tr}/D{tr})')
        for j in range(1, 14):
            c = m.cell(row=tr, column=j)
            c.border = BORDER
            c.font = BOLD
            if j > 1:
                c.number_format = PCT if j in (11, 13) else USD0
        m.cell(row=tr + 1, column=1, value='"Whole mix" = monthly totals for all companies at this scale (revenue and gross margin in $, margin as a share of revenue).').font = Font(name=F, size=9, italic=True, color='475569')
        return tr + 3

    nxt = block(1, 'Growth scale — 100 companies (plan mix on Assumptions)', 'Companies at growth scale', G['Hosting at growth scale ($ / month)'])
    block(nxt, 'Pilot scale — 5 companies (fixed hosting spread over very few companies)', 'Companies at pilot scale', G['Hosting at pilot scale ($ / month)'])
    m.freeze_panes = 'B3'
    m.cell(row=2, column=5).comment = Comment('Hosting split: a fixed share evenly per company, the rest by seats in use (Assumptions).', 'ConversaFlow')

    wb.move_sheet('Margin', offset=-1)
    wb.save(out)
    print(out)


if __name__ == '__main__':
    main(sys.argv[1])
